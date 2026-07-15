from collections import defaultdict
from datetime import timedelta

from django.db.models import Q
from django.http import HttpResponse
from django.utils import timezone
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from apps.accounts.scopes import filter_tickets_by_user
from apps.tickets.models import Ticket


def _parse_bool(value):
    return value in {"true", "True", "1", True}


def _get_date_range(request):
    today = timezone.localdate()
    start_date = today.replace(day=1)
    end_date = today

    date_from = (
        request.query_params.get("date_from")
        or request.query_params.get("created_from")
        or request.query_params.get("from")
    )
    date_to = (
        request.query_params.get("date_to")
        or request.query_params.get("created_to")
        or request.query_params.get("to")
    )

    return date_from or start_date.isoformat(), date_to or end_date.isoformat()


def _duration_minutes(ticket):
    start_at = ticket.created_at
    end_at = ticket.closed_at or ticket.done_at

    if not start_at or not end_at:
        return None

    return int((end_at - start_at).total_seconds() // 60)


def _avg(values):
    valid = [item for item in values if item is not None]

    if not valid:
        return None

    return round(sum(valid) / len(valid), 2)


def _percent(value, total):
    if not total:
        return 0

    return round((value / total) * 100, 2)


def _is_closed(ticket):
    status_code = ticket.current_status.status_code if ticket.current_status else ""

    return bool(
        ticket.closed_at
        or ticket.done_at
        or status_code in {"CLOSED", "DONE_WAIT_CLOSE"}
    )


def _display_user(user):
    if not user:
        return None

    full_name = user.get_full_name()

    return full_name or user.username or user.email


def _get_error_tickets_queryset(request):
    date_from, date_to = _get_date_range(request)

    queryset = Ticket.objects.select_related(
        "customer",
        "customer_account",
        "handling_branch",
        "assigned_unit",
        "assigned_employee",
        "owner_user",
        "current_status",
        "source",
        "priority",
        "support_category",
        "classification",
        "error_group",
        "error_type",
    ).filter(
        Q(error_group__isnull=False)
        | Q(error_type__isnull=False)
    )

    queryset = filter_tickets_by_user(queryset, request.user)

    queryset = queryset.filter(created_at__date__gte=date_from, created_at__date__lte=date_to)

    branch = request.query_params.get("branch") or request.query_params.get("handling_branch")
    source = request.query_params.get("source")
    status = request.query_params.get("status") or request.query_params.get("current_status")
    error_group = request.query_params.get("error_group")
    error_type = request.query_params.get("error_type")
    related_system = request.query_params.get("related_system")
    q = request.query_params.get("q")

    if branch:
        queryset = queryset.filter(handling_branch_id=branch)

    if source:
        queryset = queryset.filter(source_id=source)

    if status:
        if str(status).isdigit():
            queryset = queryset.filter(current_status_id=status)
        else:
            queryset = queryset.filter(current_status__status_code=status)

    if error_group:
        queryset = queryset.filter(error_group_id=error_group)

    if error_type:
        queryset = queryset.filter(error_type_id=error_type)


    if related_system:
        queryset = queryset.filter(related_system__icontains=related_system)

    if q:
        queryset = queryset.filter(
            Q(ticket_code__icontains=q)
            | Q(title__icontains=q)
            | Q(request_content__icontains=q)
            | Q(error_note__icontains=q)
            | Q(source_ref_id__icontains=q)
            | Q(error_group__group_name__icontains=q)
            | Q(error_type__type_name__icontains=q)
            | Q(customer_account__account_number__icontains=q)
        )

    return queryset.order_by("-created_at", "-id")


def _build_error_report_payload(request):
    date_from, date_to = _get_date_range(request)
    tickets = list(_get_error_tickets_queryset(request))

    total_count = len(tickets)
    closed_count = sum(1 for ticket in tickets if _is_closed(ticket))
    open_count = total_count - closed_count
    durations = [_duration_minutes(ticket) for ticket in tickets]
    avg_resolution_minutes = _avg(durations)

    category_map = {}
    source_map = {}
    status_map = {}
    branch_map = {}
    related_system_map = {}
    repeat_map = {}

    for ticket in tickets:
        duration = _duration_minutes(ticket)
        closed = _is_closed(ticket)

        group = ticket.error_group
        error_type = ticket.error_type

        category_key = (
            group.id if group else None,
            error_type.id if error_type else None,
        )

        category_item = category_map.setdefault(
            category_key,
            {
                "error_group": group.id if group else None,
                "error_group_code": group.group_code if group else None,
                "error_group_name": group.group_name if group else "Chưa phân nhóm",
                "error_type": error_type.id if error_type else None,
                "error_type_code": error_type.type_code if error_type else None,
                "error_type_name": error_type.type_name if error_type else "Chưa phân loại",
                "count": 0,
                "closed_count": 0,
                "open_count": 0,
                "duration_values": [],
            },
        )

        category_item["count"] += 1
        category_item["closed_count"] += 1 if closed else 0
        category_item["open_count"] += 0 if closed else 1
        category_item["duration_values"].append(duration)

        source = ticket.source
        source_key = source.id if source else "unknown"
        source_item = source_map.setdefault(
            source_key,
            {
                "source": source.id if source else None,
                "source_code": source.source_code if source else None,
                "source_name": source.source_name if source else "Chưa có nguồn",
                "count": 0,
                "duration_values": [],
            },
        )
        source_item["count"] += 1
        source_item["duration_values"].append(duration)

        status_obj = ticket.current_status
        status_key = status_obj.id if status_obj else "unknown"
        status_item = status_map.setdefault(
            status_key,
            {
                "status": status_obj.id if status_obj else None,
                "status_code": status_obj.status_code if status_obj else None,
                "status_name": status_obj.status_name if status_obj else "Chưa có trạng thái",
                "count": 0,
            },
        )
        status_item["count"] += 1

        branch = ticket.handling_branch
        branch_key = branch.id if branch else "unknown"
        branch_item = branch_map.setdefault(
            branch_key,
            {
                "branch": branch.id if branch else None,
                "branch_code": branch.branch_code if branch else None,
                "branch_name": branch.branch_name if branch else "Chưa có chi nhánh",
                "count": 0,
                "closed_count": 0,
                "open_count": 0,
                "duration_values": [],
            },
        )
        branch_item["count"] += 1
        branch_item["closed_count"] += 1 if closed else 0
        branch_item["open_count"] += 0 if closed else 1
        branch_item["duration_values"].append(duration)

        related_system = ticket.related_system or "Chưa xác định"
        system_item = related_system_map.setdefault(
            related_system,
            {
                "related_system": related_system,
                "count": 0,
                "duration_values": [],
            },
        )
        system_item["count"] += 1
        system_item["duration_values"].append(duration)

        repeat_key = (
            ticket.error_type_id
            or ticket.error_group_id
            or ticket.ticket_code
        )
        repeat_label = (
            ticket.error_type.type_name
            if ticket.error_type
            else ticket.error_group.group_name
            if ticket.error_group
            else ticket.ticket_code
        )
        repeat_item = repeat_map.setdefault(
            repeat_key,
            {
                "key": repeat_key,
                "label": repeat_label,
                "count": 0,
                "ticket_codes": [],
            },
        )
        repeat_item["count"] += 1
        repeat_item["ticket_codes"].append(ticket.ticket_code)

    by_category = []
    for item in category_map.values():
        item["avg_resolution_minutes"] = _avg(item.pop("duration_values"))
        item["closed_rate"] = _percent(item["closed_count"], item["count"])
        by_category.append(item)

    by_source = []
    for item in source_map.values():
        item["avg_resolution_minutes"] = _avg(item.pop("duration_values"))
        item["percentage"] = _percent(item["count"], total_count)
        by_source.append(item)

    by_status = sorted(status_map.values(), key=lambda item: item["count"], reverse=True)

    by_branch = []
    for item in branch_map.values():
        item["avg_resolution_minutes"] = _avg(item.pop("duration_values"))
        item["closed_rate"] = _percent(item["closed_count"], item["count"])
        by_branch.append(item)

    by_related_system = []
    for item in related_system_map.values():
        item["avg_resolution_minutes"] = _avg(item.pop("duration_values"))
        item["percentage"] = _percent(item["count"], total_count)
        by_related_system.append(item)

    repeated_issues = [
        item for item in repeat_map.values() if item["count"] >= 2
    ]
    repeated_ticket_count = sum(item["count"] for item in repeated_issues)

    detail = []
    for ticket in tickets:
        account_number = None

        if ticket.customer_account:
            account_number = (
                getattr(ticket.customer_account, "account_number", None)
                or getattr(ticket.customer_account, "account_no", None)
            )

        detail.append(
            {
                "id": ticket.id,
                "ticket_code": ticket.ticket_code,
                "title": ticket.title,
                "created_at": ticket.created_at,
                "updated_at": ticket.updated_at,
                "branch_name": ticket.handling_branch.branch_name if ticket.handling_branch else None,
                "source_name": ticket.source.source_name if ticket.source else None,
                "status_name": ticket.current_status.status_name if ticket.current_status else None,
                "priority_name": ticket.priority.priority_name if ticket.priority else None,
                "error_group_name": ticket.error_group.group_name if ticket.error_group else None,
                "error_type_name": ticket.error_type.type_name if ticket.error_type else None,
                "error_note": ticket.error_note,
                "related_system": ticket.related_system,
                "source_ref_id": ticket.source_ref_id,
                "external_status": ticket.external_status,
                "account_number": account_number,
                "customer_name": ticket.customer.full_name if ticket.customer else None,
                "assigned_employee_name": ticket.assigned_employee.full_name if ticket.assigned_employee else None,
                "owner_user_name": _display_user(ticket.owner_user),
                "duration_minutes": _duration_minutes(ticket),
                "is_closed": _is_closed(ticket),
            }
        )

    return {
        "filters": {
            "date_from": date_from,
            "date_to": date_to,
        },
        "overview": {
            "total_errors": total_count,
            "open_errors": open_count,
            "closed_errors": closed_count,
            "avg_resolution_minutes": avg_resolution_minutes,
            "repeated_issue_count": len(repeated_issues),
            "repeated_ticket_count": repeated_ticket_count,
            "repeated_rate": _percent(repeated_ticket_count, total_count),
        },
        "by_category": sorted(by_category, key=lambda item: item["count"], reverse=True),
        "by_source": sorted(by_source, key=lambda item: item["count"], reverse=True),
        "by_status": by_status,
        "by_branch": sorted(by_branch, key=lambda item: item["count"], reverse=True),
        "by_related_system": sorted(by_related_system, key=lambda item: item["count"], reverse=True),
        "repeated_issues": sorted(repeated_issues, key=lambda item: item["count"], reverse=True),
        "detail": detail,
    }


class TicketErrorReportAPIView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        return Response(_build_error_report_payload(request))


def _write_rows(sheet, rows):
    for row in rows:
        sheet.append(row)


class TicketErrorReportExportAPIView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        try:
            from openpyxl import Workbook
        except ImportError as exc:
            raise RuntimeError(
                "Cần cài openpyxl để xuất Excel: pip install openpyxl"
            ) from exc

        payload = _build_error_report_payload(request)
        workbook = Workbook()

        sheet = workbook.active
        sheet.title = "Theo danh mục"
        _write_rows(
            sheet,
            [
                [
                    "Nhóm lỗi",
                    "Loại lỗi",
                    "Số lượng",
                    "Đã xử lý",
                    "Đang mở",
                    "Tỷ lệ xử lý (%)",
                    "Thời gian xử lý TB (phút)",
                ]
            ],
        )
        for item in payload["by_category"]:
            sheet.append(
                [
                    item["error_group_name"],
                    item["error_type_name"],
                    item["count"],
                    item["closed_count"],
                    item["open_count"],
                    item["closed_rate"],
                    item["avg_resolution_minutes"],
                ]
            )

        source_sheet = workbook.create_sheet("Theo nguồn")
        _write_rows(
            source_sheet,
            [
                [
                    "Kênh tiếp nhận",
                    "Số lượng lỗi",
                    "Tỷ trọng (%)",
                    "Thời gian xử lý TB (phút)",
                ]
            ],
        )
        for item in payload["by_source"]:
            source_sheet.append(
                [
                    item["source_name"],
                    item["count"],
                    item["percentage"],
                    item["avg_resolution_minutes"],
                ]
            )

        detail_sheet = workbook.create_sheet("Chi tiết")
        _write_rows(
            detail_sheet,
            [
                [
                    "Mã ticket",
                    "Ngày tạo",
                    "Chi nhánh",
                    "Kênh",
                    "Trạng thái",
                    "Nhóm lỗi",
                    "Loại lỗi",
                    "Hệ thống liên quan",
                    "Mã tham chiếu ngoài",
                    "Số tài khoản",
                    "Khách hàng",
                    "Người xử lý",
                    "Thời gian xử lý (phút)",
                    "Ghi chú lỗi",
                ]
            ],
        )

        for item in payload["detail"]:
            detail_sheet.append(
                [
                    item["ticket_code"],
                    item["created_at"],
                    item["branch_name"],
                    item["source_name"],
                    item["status_name"],
                    item["error_group_name"],
                    item["error_type_name"],
                    item["related_system"],
                    item["source_ref_id"],
                    item["account_number"],
                    item["customer_name"],
                    item["assigned_employee_name"],
                    item["duration_minutes"],
                    item["error_note"],
                ]
            )

        today_label = timezone.localdate().strftime("%Y%m%d")
        filename = f"bao_cao_loi_{today_label}.xlsx"

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        workbook.save(response)

        return response
