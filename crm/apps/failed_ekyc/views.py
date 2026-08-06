import io
import re
import unicodedata
from collections import Counter
from datetime import date, datetime, timedelta

from django.db import transaction
from django.db.models import Count, Q
from django.http import HttpResponse
from rest_framework import status, viewsets
from rest_framework.decorators import action
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from apps.customers.models import Customer, CustomerAccount
from apps.failed_ekyc.models import FailedEkycRecord
from apps.failed_ekyc.permissions import HasFailedEkycPermission
from apps.failed_ekyc.serializers import FailedEkycRecordSerializer


def normalize(value):
    value = unicodedata.normalize("NFD", str(value or "").strip().casefold())
    return " ".join("".join(c for c in value if unicodedata.category(c) != "Mn").replace("đ", "d").split())


def parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


class FailedEkycRecordViewSet(viewsets.ModelViewSet):
    permission_classes = [IsAuthenticated, HasFailedEkycPermission]
    serializer_class = FailedEkycRecordSerializer
    queryset = FailedEkycRecord.objects.all()

    def get_queryset(self):
        qs = FailedEkycRecord.objects.select_related("customer", "customer_account", "created_by_user")
        p = self.request.query_params
        q = p.get("q")
        if q:
            qs = qs.filter(Q(account_number__icontains=q) | Q(customer_name__icontains=q) | Q(email__icontains=q) | Q(phone__icontains=q) | Q(error_message__icontains=q) | Q(cs_comment__icontains=q))
        for field in ("step", "call_status", "call_result"):
            if p.get(field):
                qs = qs.filter(**{field: p[field]})
        for field in ("account_number", "customer_name", "email", "phone", "branch_name", "pic", "error_message", "cs_comment"):
            if p.get(field):
                qs = qs.filter(**{f"{field}__icontains": p[field]})
        if p.get("follow_count"):
            try:
                qs = qs.filter(follow_count=int(p["follow_count"].replace("+", "")))
            except (ValueError, AttributeError):
                pass
        for field, param in (("failed_at", "failed_date"), ("call_date", "call_date")):
            if p.get(f"{param}_from"):
                qs = qs.filter(**{f"{field}__gte": p[f"{param}_from"]})
            if p.get(f"{param}_to"):
                qs = qs.filter(**{f"{field}__lte": p[f"{param}_to"]})
        return qs.order_by("-failed_at", "-id")

    def perform_create(self, serializer):
        serializer.save(created_by_user=self.request.user)

    @action(detail=False, methods=["get"], url_path="lookup-customer")
    def lookup_customer(self, request):
        account_number = request.query_params.get("account_number", "").strip().upper()
        if not account_number:
            return Response({"found": False, "message": "Vui lòng nhập số tài khoản."}, status=400)
        account = CustomerAccount.objects.select_related("customer", "customer__branch").filter(account_number__iexact=account_number).first()
        if not account:
            return Response({"found": False, "account_number": account_number}, status=404)
        customer = account.customer
        return Response({
            "found": True,
            "customer_id": customer.id if customer else None,
            "customer_account_id": account.id,
            "account_number": account.account_number,
            "customer_name": customer.full_name if customer else "",
            "branch_name": customer.branch.branch_name if customer and customer.branch else "",
            "email": getattr(customer, "email", "") if customer else "",
            "phone": customer.phone if customer else "",
        })

    @action(detail=False, methods=["get"])
    def dashboard(self, request):
        today = date.today()
        start_date = parse_date(request.query_params.get("failed_date_from")) or (today - timedelta(days=150))
        end_date = parse_date(request.query_params.get("failed_date_to")) or today
        granularity = request.query_params.get("granularity", "MONTH").upper()
        compare_mode = request.query_params.get("compare_mode", "NONE").upper()
        qs = FailedEkycRecord.objects.select_related("customer", "customer_account", "created_by_user").filter(
            Q(failed_at__gte=start_date, failed_at__lte=end_date) | Q(failed_at__isnull=True)
        )
        status_rows = list(qs.values("call_status").annotate(count=Count("id")).order_by("call_status"))
        result_rows = list(qs.values("call_result").annotate(count=Count("id")).order_by("call_result"))
        step_rows = list(qs.values("step").annotate(count=Count("id")).order_by("step"))

        def period_label(value):
            if granularity == "YEAR":
                return value.strftime("%Y")
            if granularity == "QUARTER":
                return f"Q{((value.month - 1) // 3) + 1}/{value.year}"
            return value.strftime("%m/%Y")

        trend_map = {}
        for record in qs.order_by("failed_at", "id"):
            effective_date = record.failed_at or end_date
            label = period_label(effective_date)
            item = trend_map.setdefault(label, {
                "date": label, "sort_date": effective_date, "count": 0, "called": 0,
                "not_called": 0, "success": 0, "retry_required": 0, "failed": 0,
                "contact_success": 0, "contact_failed": 0,
                "guidance_success": 0, "guidance_failed": 0, "retry_later": 0, "system_test": 0,
                "no_answer": 0, "hung_up": 0, "invalid_number": 0, "call_back": 0, "blocked_number": 0,
                "customer_completed": 0, "customer_no_action": 0,
            })
            item["count"] += 1
            status_text = normalize(record.call_status)
            detail_text = normalize(f"{record.call_status} {record.call_result} {record.cs_comment}")
            is_not_called = not record.call_status or status_text in ("khong call", "khong goi")
            is_answered = status_text == "nghe may"
            if not is_not_called:
                item["called"] += 1
            else:
                item["not_called"] += 1
            if is_answered:
                item["contact_success"] += 1
            elif not is_not_called:
                item["contact_failed"] += 1
            if record.call_result == "Thành công":
                item["success"] += 1
            elif record.call_result == "KH cần thử lại":
                item["retry_required"] += 1
            elif record.call_result == "Không thành công":
                item["failed"] += 1
            if is_answered:
                if "test ht" in detail_text or "test he thong" in detail_text:
                    item["system_test"] += 1
                elif record.call_result == "Thành công":
                    item["guidance_success"] += 1
                    item["customer_completed"] += 1
                elif record.call_result == "KH cần thử lại":
                    item["retry_later"] += 1
                    item["customer_no_action"] += 1
                elif record.call_result == "Không thành công":
                    item["guidance_failed"] += 1
                    item["customer_no_action"] += 1
                else:
                    item["customer_no_action"] += 1
            elif not is_not_called:
                if "tat may" in detail_text or "may ngang" in detail_text:
                    item["hung_up"] += 1
                elif "thue bao" in detail_text or "khong ton tai" in detail_text:
                    item["invalid_number"] += 1
                elif "ban" in detail_text or "goi lai" in detail_text:
                    item["call_back"] += 1
                elif "chan so" in detail_text or "khoa may" in detail_text:
                    item["blocked_number"] += 1
                else:
                    item["no_answer"] += 1
        trends = []
        for item in sorted(trend_map.values(), key=lambda value: value["sort_date"]):
            item.pop("sort_date")
            trends.append(item)

        comparison = None
        if compare_mode in ("YOY", "QOQ"):
            span = max(1, (end_date - start_date).days)
            if compare_mode == "YOY":
                prev_start, prev_end = start_date - timedelta(days=365), end_date - timedelta(days=365)
            else:
                prev_end = start_date - timedelta(days=1)
                prev_start = prev_end - timedelta(days=span)
            prev_total = FailedEkycRecord.objects.filter(failed_at__gte=prev_start, failed_at__lte=prev_end).count()
            current_total = qs.count()
            growth = round(((current_total - prev_total) / prev_total) * 100, 1) if prev_total else (100.0 if current_total else 0.0)
            comparison = {"compare_mode": compare_mode, "prev_total": prev_total, "total_growth_percent": growth, "prev_start_date": prev_start, "prev_end_date": prev_end}

        recent_records = FailedEkycRecordSerializer(qs.order_by("-failed_at", "-id")[:10], many=True).data
        no_call_records = list(qs.filter(Q(call_status="KHÔNG CALL") | Q(call_status="") | Q(call_status__isnull=True)))
        duplicate_keys = Counter(
            (record.step, record.account_number, record.failed_at)
            for record in no_call_records if record.account_number
        )
        already_has_account = sum(
            1 for record in no_call_records
            if any(keyword in normalize(record.cs_comment) for keyword in ("da mtk", "da mo tai khoan", "da co tk"))
        )
        duplicate_cases = sum(max(0, count - 1) for count in duplicate_keys.values())
        error_rules = [
            ("Khuôn mặt không khớp", ("khuon mat khong khop",)),
            ("Nhận diện người thật thất bại", ("nhan dien nguoi that", "liveness")),
            ("Giấy tờ không hợp lệ", ("giay to khong hop le", "gttt khong hop le")),
            ("Ảnh mờ/bóng", ("mo/bong", "bi mo", "bi bong")),
            ("Ảnh chụp quá xa", ("chup qua xa",)),
            ("Ảnh bị che/mất góc", ("bi che", "mat goc", "sat canh")),
            ("Nhiều giấy tờ/chữ ký", ("nhieu hon mot giay to", "nhieu hon mot chu ky")),
            ("Thiếu/sai chữ ký", ("chua co chu ky", "thieu chu ky", "chua ghi ro ho va ten")),
            ("Dịch vụ bảo trì", ("bao tri",)),
        ]
        error_counts = Counter()
        for record in qs:
            error_text = normalize(record.error_message)
            matched = False
            for label, keywords in error_rules:
                if any(keyword in error_text for keyword in keywords):
                    error_counts[label] += 1
                    matched = True
            if error_text and not matched:
                error_counts["Lỗi khác"] += 1

        pic_map = {}
        for record in qs.exclude(pic=""):
            pic_name = record.pic.strip()
            if not pic_name:
                continue
            item = pic_map.setdefault(pic_name, {"pic": pic_name, "total": 0, "contacted": 0, "successful": 0})
            item["total"] += 1
            if record.call_status == "Nghe máy":
                item["contacted"] += 1
            if record.call_result == "Thành công":
                item["successful"] += 1
        pic_performance = sorted(pic_map.values(), key=lambda item: (-item["total"], item["pic"]))[:10]
        return Response({
            "total_records": qs.count(), "granularity": granularity, "compare_mode": compare_mode,
            "start_date": start_date, "end_date": end_date, "status_breakdown": status_rows,
            "result_breakdown": result_rows, "step_breakdown": step_rows, "daily_trends": trends,
            "comparison": comparison, "recent_records": recent_records,
            "no_call_details": [
                {"category": "Đã có TK tại PHS", "count": already_has_account},
                {"category": "Trùng Case", "count": duplicate_cases},
            ],
            "top_errors": [
                {"error": label, "count": count}
                for label, count in error_counts.most_common(8)
            ],
            "pic_performance": pic_performance,
        })

    @action(detail=False, methods=["post"], url_path="import-excel")
    def import_excel(self, request):
        file_obj = request.FILES.get("file")
        if not file_obj:
            return Response({"detail": "Vui lòng đính kèm file Excel (.xlsx)."}, status=400)
        if file_obj.size > 10 * 1024 * 1024:
            return Response({"detail": "File vượt quá giới hạn 10 MB."}, status=400)
        try:
            from openpyxl import load_workbook
            rows = list(load_workbook(file_obj, data_only=True).active.iter_rows(values_only=True))
        except Exception as exc:
            return Response({"detail": f"Không thể đọc file Excel: {exc}"}, status=400)
        if len(rows) < 2:
            return Response({"detail": "File không có dữ liệu."}, status=400)
        aliases = {
            "step": ("step", "buoc"), "branch_name": ("branch", "chi nhanh"), "account_number": ("account", "so tai khoan", "so tk"),
            "email": ("email",), "phone": ("phone number", "phone", "so dien thoai", "sdt"), "failed_at": ("created date", "ngay tao"),
            "error_message": ("error", "loi"), "pic": ("pic",), "call_date": ("ngay goi",), "follow_count": ("so lan follow", "follow"),
            "call_status": ("tinh trang",), "call_result": ("ket qua cuoc goi", "ket qua"), "cs_comment": ("cs comment", "ghi chu"),
        }
        headers = [normalize(v) for v in rows[0]]
        columns = {}
        for field, names in aliases.items():
            for i, header in enumerate(headers):
                if any(normalize(name) in header for name in names):
                    columns[field] = i
                    break
        def cell(row, field):
            index = columns.get(field)
            return row[index] if index is not None and index < len(row) else None

        records, errors = [], []
        valid_steps = {value for value, _ in FailedEkycRecord.STEP_CHOICES}
        for row_number, row in enumerate(rows[1:], 2):
            if not any(value not in (None, "") for value in row):
                continue
            account_number = str(cell(row, "account_number") or "").strip().upper()
            email = str(cell(row, "email") or "").strip()
            phone = str(cell(row, "phone") or "").strip()
            try:
                step = str(cell(row, "step") or "EKYC").strip().upper()
                if step not in valid_steps:
                    errors.append(f"Dòng {row_number}: Step không hợp lệ ({step})")
                    continue
                follow_text = str(cell(row, "follow_count") or "")
                match = re.search(r"\d+", follow_text)
                follow_count = int(match.group()) if match else 0
                call_status = str(cell(row, "call_status") or "").strip()
                if normalize(call_status) in ("khong call", "khong goi"):
                    call_status = "KHÔNG CALL"
                account = CustomerAccount.objects.select_related("customer").filter(account_number__iexact=account_number).first() if account_number else None
                customer = account.customer if account else None
                records.append(FailedEkycRecord(
                    step=step, branch_name=str(cell(row, "branch_name") or "").strip(), customer=customer, customer_account=account,
                    account_number=account_number, customer_name=customer.full_name if customer else "", email=email, phone=phone or (customer.phone if customer else ""),
                    failed_at=parse_date(cell(row, "failed_at")), error_message=str(cell(row, "error_message") or "").strip(), pic=str(cell(row, "pic") or "").strip(),
                    call_date=parse_date(cell(row, "call_date")), follow_count=follow_count, call_status=call_status,
                    call_result=str(cell(row, "call_result") or "").strip(), cs_comment=str(cell(row, "cs_comment") or "").strip(), created_by_user=request.user,
                ))
            except Exception as exc:
                errors.append(f"Dòng {row_number}: {exc}")
        with transaction.atomic():
            FailedEkycRecord.objects.bulk_create(records, batch_size=500)
        return Response({"message": "Import hoàn tất.", "success_count": len(records), "error_count": len(errors), "errors": errors})

    @action(detail=False, methods=["get"], url_path="download-template")
    def download_template(self, request):
        return self._excel_response([], "Failed_eKYC_Import_Template.xlsx", include_sample=True)

    @action(detail=False, methods=["get"], url_path="export-excel")
    def export_excel(self, request):
        return self._excel_response(self.filter_queryset(self.get_queryset()), "Failed_eKYC_Export.xlsx")

    def _excel_response(self, records, filename, include_sample=False):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Failed eKYC"
        headers = ["Step", "Branch", "Account", "Email", "Phone number", "Created date", "Error", "PIC", "Ngày gọi", "Số lần follow", "Tình Trạng", "Kết quả cuộc gọi", "CS comment"]
        ws.append(headers)
        if include_sample:
            ws.append(["EKYC", "", "022C119340", "", "", "04/07/2026", "Mặt trước giấy tờ bị che", "Mai", "06/07/2026", "Lần 2", "Nghe máy", "Thành công", "KH đã mở tài khoản thành công."])
        for item in records:
            ws.append([item.step, item.branch_name, item.account_number, item.email, item.phone, item.failed_at, item.error_message, item.pic, item.call_date, item.follow_count, item.call_status, item.call_result, item.cs_comment])
        output = io.BytesIO()
        wb.save(output)
        response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response
