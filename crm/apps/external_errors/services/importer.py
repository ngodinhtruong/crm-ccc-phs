
import re
import unicodedata
import uuid
from datetime import date, datetime, time
from typing import Any

from django.db import transaction
from django.db.models import Q
from django.utils import timezone

from apps.external_errors.models import (
    ExternalErrorImportBatch,
    ExternalErrorRecord,
)
from apps.external_errors.services.cleaning import (
    build_rule_based_clean_fields,
    clean_text,
)
from apps.external_errors.services.dashboard_cache import (
    invalidate_external_error_dashboard_cache,
)


EXCEL_COLUMN_ALIASES = {
    "received_date": ("Ngày nhận",),
    "completed_date": ("Ngày hoàn thành",),
    "source": ("Nguồn",),
    "device": ("Thiết bị",),
    "result": ("Kết quả xử lý",),
    "content": ("Nội dung",),
    "cause": (
        "Nguyên nhân",
        "Nguyên nhân lỗi",
        "Nguyên nhân chính",
    ),
    "solution": (
        "Giải pháp",
        "Giải pháp xử lý",
        "Biện pháp khắc phục",
    ),
}

OPTIONAL_EXCEL_FIELDS = {"solution"}

# Alias tương thích với code/test cũ.
EXCEL_COLUMNS = {
    field_name: aliases[0]
    for field_name, aliases in EXCEL_COLUMN_ALIASES.items()
}


def normalize_header(value: Any) -> str:
    text = clean_text(value).lower()
    text = unicodedata.normalize("NFD", text)
    text = "".join(
        char
        for char in text
        if unicodedata.category(char) != "Mn"
    )
    return re.sub(r"\s+", " ", text).strip()


def ensure_aware(value: datetime) -> datetime:
    if timezone.is_aware(value):
        return value
    return timezone.make_aware(
        value,
        timezone.get_current_timezone(),
    )


def normalize_datetime_string(text: Any) -> str:
    cleaned = clean_text(text)
    if not cleaned:
        return ""

    # Gom nhiều khoảng trắng dư thừa thành 1 khoảng trắng duy nhất
    cleaned = re.sub(r"\s+", " ", cleaned).strip()

    # Xóa khoảng trắng thừa xung quanh các ký tự phân cách ngày giờ như /, -, :
    # Ví dụ: "06 / 05 / 2026   13 : 54 : 00" -> "06/05/2026 13:54:00"
    cleaned = re.sub(r"\s*([/\-:])\s*", r"\1", cleaned)

    # Gom lại khoảng trắng nếu có sau khi làm sạch ký tự phân cách
    cleaned = re.sub(r"\s+", " ", cleaned).strip()

    def fix_time_ampm(m: re.Match) -> str:
        time_part = m.group(1)
        ampm = m.group(2).upper()

        parts = time_part.split(":")
        try:
            hour = int(parts[0])
        except ValueError:
            return f"{time_part} {ampm}"

        if hour > 12 or hour == 0:
            # Ví dụ: 13:54 PM / 13:54 AM -> 13:54, 00:15 PM -> 00:15
            return time_part
        elif hour == 12:
            if ampm == "AM":
                # 12:30 AM -> 00:30
                parts[0] = "00"
                return ":".join(parts)
            else:
                # 12:30 PM -> 12:30
                return time_part
        else:
            # 1 <= hour < 12 (ví dụ: 02:30 PM -> 02:30 PM)
            return f"{time_part} {ampm}"

    return re.sub(
        r"(?i)(\d{1,2}:\d{2}(?::\d{2})?)\s*(am|pm)\b",
        fix_time_ampm,
        cleaned,
    )


COMMON_DATETIME_FORMATS = (
    # 12-hour AM/PM formats
    "%d/%m/%Y %I:%M:%S %p",
    "%d/%m/%Y %I:%M %p",
    "%d-%m-%Y %I:%M:%S %p",
    "%d-%m-%Y %I:%M %p",
    "%Y-%m-%d %I:%M:%S %p",
    "%Y-%m-%d %I:%M %p",
    "%m/%d/%Y %I:%M:%S %p",
    "%m/%d/%Y %I:%M %p",
    "%m-%d-%Y %I:%M:%S %p",
    "%m-%d-%Y %I:%M %p",
    # 24-hour formats
    "%d/%m/%Y %H:%M:%S",
    "%d/%m/%Y %H:%M",
    "%d/%m/%Y",
    "%d-%m-%Y %H:%M:%S",
    "%d-%m-%Y %H:%M",
    "%d-%m-%Y",
    "%Y-%m-%dT%H:%M:%S",
    "%Y-%m-%dT%H:%M",
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%d %H:%M",
    "%Y-%m-%d",
    "%m/%d/%Y %H:%M:%S",
    "%m/%d/%Y %H:%M",
    "%m/%d/%Y",
    "%m-%d-%Y %H:%M:%S",
    "%m-%d-%Y %H:%M",
    "%m-%d-%Y",
)


def parse_received_datetime(value: Any) -> datetime | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return ensure_aware(value)
    if isinstance(value, date):
        return ensure_aware(datetime.combine(value, time.min))

    text_value = normalize_datetime_string(value)
    if not text_value:
        return None

    for fmt in COMMON_DATETIME_FORMATS:
        try:
            return ensure_aware(
                datetime.strptime(text_value, fmt)
            )
        except ValueError:
            continue
    return None


def parse_completed_datetime(value: Any) -> datetime | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        if value.time() == time.min:
            value = datetime.combine(
                value.date(),
                time(23, 59),
            )
        return ensure_aware(value)
    if isinstance(value, date):
        return ensure_aware(
            datetime.combine(value, time(23, 59))
        )

    text_value = normalize_datetime_string(value)
    if not text_value:
        return None

    has_time = bool(
        re.search(
            r"(?i)\b\d{1,2}:\d{2}(?::\d{2})?\b|\b(am|pm)\b",
            text_value,
        )
    )

    for fmt in COMMON_DATETIME_FORMATS:
        try:
            parsed = datetime.strptime(text_value, fmt)
            if not has_time:
                parsed = datetime.combine(
                    parsed.date(),
                    time(23, 59),
                )
            return ensure_aware(parsed)
        except ValueError:
            continue
    return None


def is_resolved_result(value: Any) -> bool:
    text = clean_text(value)
    if not text:
        return False
    text = unicodedata.normalize("NFD", text.lower())
    text = "".join(
        char
        for char in text
        if unicodedata.category(char) != "Mn"
    )
    text = re.sub(r"[^a-z0-9]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return bool(
        re.search(
            r"\bda(?: duoc)? khac phuc\b|\bhoan thanh\b|\bda xu ly\b|\bresolved\b|\bdone\b|\bfixed\b",
            text,
        )
    )


def build_batch_code(prefix: str = "EXTERR") -> str:
    timestamp = timezone.localtime().strftime(
        "%Y%m%d%H%M%S"
    )
    return (
        f"{prefix}-{timestamp}-"
        f"{uuid.uuid4().hex[:6].upper()}"
    )


def build_record(
    *,
    received_date: datetime,
    completed_date: datetime | None = None,
    source: Any,
    device: Any,
    result: Any,
    content: Any,
    cause: Any = None,
    solution: Any = None,
    batch=None,
    created_by=None,
) -> ExternalErrorRecord:
    if received_date is None:
        raise ValueError(
            "Không thể tạo record khi Ngày nhận bị rỗng."
        )

    raw_data = {
        "raw_source": clean_text(source) or None,
        "raw_device": clean_text(device) or None,
        "raw_result": clean_text(result) or None,
        "raw_content": clean_text(content) or None,
        "raw_cause": clean_text(cause) or None,
        "raw_solution": clean_text(solution) or None,
    }
    clean_fields = build_rule_based_clean_fields(raw_data)

    return ExternalErrorRecord(
        batch=batch,
        received_date=received_date,
        completed_date=completed_date,
        **raw_data,
        **clean_fields,
        error_code=None,
        normalized_issue=None,
        classification_confidence=None,
        need_review=False,
        classification_status=(
            ExternalErrorRecord.STATUS_UNCLASSIFIED
        ),
        classification_error=None,
        llm_model_id=None,
        classified_at=None,
        cause_group=None,
        normalized_cause=None,
        cause_classification_confidence=None,
        cause_need_review=False,
        cause_classification_status=(
            ExternalErrorRecord.STATUS_UNCLASSIFIED
        ),
        cause_classification_error=None,
        cause_llm_model_id=None,
        cause_classified_at=None,
        created_by=created_by,
        updated_by=created_by,
    )


def validate_record_values(
    *,
    received_date: Any,
    completed_date: Any,
    result: Any,
    content: Any,
    import_time: datetime | None = None,
) -> tuple[datetime, datetime | None]:
    received_at = parse_received_datetime(received_date)
    if received_at is None:
        raise ValueError(
            "Ngày nhận bị thiếu hoặc không đúng định dạng."
        )

    if not clean_text(content):
        raise ValueError("Nội dung không được để trống.")

    completed_at = parse_completed_datetime(completed_date)
    if completed_at is None:
        if is_resolved_result(result):
            completed_at = import_time or timezone.now()
        else:
            completed_at = None

    if completed_at is not None and completed_at < received_at:
        completed_at = received_at

    return received_at, completed_at


def create_manual_record(
    *,
    data: dict,
    created_by,
) -> ExternalErrorRecord:
    received_at, completed_at = validate_record_values(
        received_date=data.get("received_date"),
        completed_date=data.get("completed_date"),
        result=data.get("result"),
        content=data.get("content"),
        import_time=timezone.now(),
    )

    record = build_record(
        received_date=received_at,
        completed_date=completed_at,
        source=data.get("source"),
        device=data.get("device"),
        result=data.get("result"),
        content=data.get("content"),
        cause=data.get("cause"),
        solution=data.get("solution"),
        created_by=created_by,
    )
    record.save()
    try:
        from apps.external_errors.models import ExternalErrorRecordAuditLog
        from apps.external_errors.services.audit import log_record_audit, snapshot_record_data
        log_record_audit(
            record=record,
            action_type=ExternalErrorRecordAuditLog.ACTION_CREATE,
            user=created_by,
            new_data=snapshot_record_data(record),
            note="Tạo mới lỗi thủ công",
        )
    except Exception:
        pass
    invalidate_external_error_dashboard_cache()
    return record


def find_excel_columns(
    headers: tuple[Any, ...],
) -> dict[str, int | None]:
    normalized_headers = {
        normalize_header(header): index
        for index, header in enumerate(headers)
        if clean_text(header)
    }

    result = {}
    missing = []

    for field_name, aliases in EXCEL_COLUMN_ALIASES.items():
        column_index = None
        for alias in aliases:
            column_index = normalized_headers.get(
                normalize_header(alias)
            )
            if column_index is not None:
                break

        if column_index is None:
            if field_name in OPTIONAL_EXCEL_FIELDS:
                result[field_name] = None
            else:
                missing.append(aliases[0])
        else:
            result[field_name] = column_index

    if missing:
        raise ValueError(
            "File Excel thiếu các cột bắt buộc: "
            + ", ".join(missing)
        )

    return result


def read_excel_rows(
    uploaded_file,
    sheet_name: str | None = None,
):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError(
            "Thiếu openpyxl. Chạy: pip install openpyxl"
        ) from exc

    if hasattr(uploaded_file, "seek"):
        uploaded_file.seek(0)

    workbook = load_workbook(
        filename=uploaded_file,
        read_only=True,
        data_only=True,
    )

    if sheet_name:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"Không tìm thấy sheet '{sheet_name}'. "
                "Các sheet hiện có: "
                + ", ".join(workbook.sheetnames)
            )
        worksheet = workbook[sheet_name]
    else:
        worksheet = workbook[workbook.sheetnames[0]]

    row_iterator = worksheet.iter_rows(values_only=True)
    try:
        headers = next(row_iterator)
    except StopIteration as exc:
        raise ValueError(
            "Sheet Excel không có dữ liệu."
        ) from exc

    return (
        worksheet.title,
        find_excel_columns(headers),
        row_iterator,
    )


def classify_batch_records(batch, records):
    from apps.external_errors.services.bedrock_classifier import (
        classify_queryset,
    )
    from apps.external_errors.services.bedrock_cause_classifier import (
        classify_queryset_causes,
    )

    batch.status = ExternalErrorImportBatch.STATUS_CLASSIFYING
    batch.save(update_fields=["status", "updated_at"])

    queryset = batch.records.all().order_by("id")
    error_stats = classify_queryset(queryset, force=False)
    cause_stats = classify_queryset_causes(
        queryset,
        force=False,
    )

    successful_statuses = [
        ExternalErrorRecord.STATUS_CLASSIFIED,
        ExternalErrorRecord.STATUS_NEED_REVIEW,
        ExternalErrorRecord.STATUS_CONFIRMED,
    ]
    fully_processed = batch.records.filter(
        classification_status__in=successful_statuses,
        cause_classification_status__in=successful_statuses,
    ).count()
    failed_records = batch.records.filter(
        Q(
            classification_status=(
                ExternalErrorRecord.STATUS_FAILED
            )
        )
        | Q(
            cause_classification_status=(
                ExternalErrorRecord.STATUS_FAILED
            )
        )
    ).distinct().count()

    batch.classified_rows = fully_processed
    batch.failed_rows += failed_records
    batch.status = (
        ExternalErrorImportBatch.STATUS_FAILED
        if fully_processed == 0 and failed_records > 0
        else ExternalErrorImportBatch.STATUS_CLASSIFIED
    )
    batch.save(
        update_fields=[
            "classified_rows",
            "failed_rows",
            "status",
            "updated_at",
        ]
    )

    invalidate_external_error_dashboard_cache()

    return {
        "error_classification": error_stats,
        "cause_classification": cause_stats,
    }


def import_excel_file(
    *,
    uploaded_file,
    sheet_name: str | None,
    auto_classify: bool,
    created_by,
):
    import_time = timezone.now()
    file_name = (
        getattr(uploaded_file, "name", "")
        or "external-errors.xlsx"
    )
    sheet, column_map, rows = read_excel_rows(
        uploaded_file,
        sheet_name,
    )

    pending_records = []
    skipped = []
    source_rows = 0

    for row_number, row in enumerate(rows, start=2):
        values = {
            field_name: (
                row[column_index]
                if column_index is not None and column_index < len(row)
                else None
            )
            for field_name, column_index in column_map.items()
        }

        if not any(
            clean_text(value)
            for value in values.values()
        ):
            continue

        source_rows += 1

        try:
            received_at, completed_at = validate_record_values(
                received_date=values["received_date"],
                completed_date=values["completed_date"],
                result=values["result"],
                content=values["content"],
                import_time=import_time,
            )
            pending_records.append(
                build_record(
                    received_date=received_at,
                    completed_date=completed_at,
                    source=values["source"],
                    device=values["device"],
                    result=values["result"],
                    content=values["content"],
                    cause=values.get("cause"),
                    solution=values.get("solution"),
                    created_by=created_by,
                )
            )
        except ValueError as exc:
            skipped.append(
                {
                    "row": row_number,
                    "reason": str(exc),
                }
            )

    with transaction.atomic():
        batch = ExternalErrorImportBatch.objects.create(
            batch_code=build_batch_code(),
            file_name=file_name,
            source_type=ExternalErrorImportBatch.SOURCE_EXCEL,
            status=ExternalErrorImportBatch.STATUS_IMPORTED,
            created_by=created_by,
        )

        for record in pending_records:
            record.batch = batch

        ExternalErrorRecord.objects.bulk_create(
            pending_records,
            batch_size=500,
        )

        invalid_count = batch.records.filter(
            received_date__isnull=True
        ).count()
        if invalid_count:
            raise RuntimeError(
                "Import bị hủy vì có "
                f"{invalid_count} record không lưu được Ngày nhận."
            )

        batch.total_rows = len(pending_records)
        batch.failed_rows = len(skipped)
        batch.save(
            update_fields=[
                "total_rows",
                "failed_rows",
                "updated_at",
            ]
        )
        transaction.on_commit(
            invalidate_external_error_dashboard_cache
        )

    classification_stats = {
        "error_classification": {
            "total": 0,
            "classified": 0,
            "failed": 0,
            "errors": [],
        },
        "cause_classification": {
            "total": 0,
            "classified": 0,
            "need_review": 0,
            "failed": 0,
            "errors": [],
        },
    }

    if auto_classify and pending_records:
        classification_stats = classify_batch_records(
            batch,
            pending_records,
        )

    return batch, {
        "sheet_name": sheet,
        "source_rows": source_rows,
        "imported_rows": len(pending_records),
        "skipped_rows": len(skipped),
        "skipped_details": skipped,
        "auto_classify": auto_classify,
        "classification": classification_stats,
    }


def import_and_optionally_classify(
    *,
    rows,
    file_name="",
    source_type=ExternalErrorImportBatch.SOURCE_API,
    classify_now=True,
    created_by=None,
):
    batch = ExternalErrorImportBatch.objects.create(
        batch_code=build_batch_code("EXTERR-API"),
        file_name=file_name,
        source_type=source_type,
        status=ExternalErrorImportBatch.STATUS_IMPORTED,
        created_by=created_by,
    )
    import_time = timezone.now()
    records = []
    skipped = []

    for index, row in enumerate(rows, start=1):
        try:
            received_at, completed_at = validate_record_values(
                received_date=row.get("received_date"),
                completed_date=row.get("completed_date"),
                result=row.get("result"),
                content=row.get("content"),
                import_time=import_time,
            )
        except ValueError as exc:
            skipped.append(
                {"row": index, "reason": str(exc)}
            )
            continue

        records.append(
            build_record(
                batch=batch,
                received_date=received_at,
                completed_date=completed_at,
                source=row.get("source"),
                device=row.get("device"),
                result=row.get("result"),
                content=row.get("content"),
                cause=row.get("cause"),
                solution=row.get("solution"),
                created_by=created_by,
            )
        )

    with transaction.atomic():
        ExternalErrorRecord.objects.bulk_create(
            records,
            batch_size=500,
        )
        batch.total_rows = len(records)
        batch.failed_rows = len(skipped)
        batch.save(
            update_fields=[
                "total_rows",
                "failed_rows",
                "updated_at",
            ]
        )
        transaction.on_commit(
            invalidate_external_error_dashboard_cache
        )

    if classify_now and records:
        classify_batch_records(batch, records)

    return batch
